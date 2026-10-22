import json
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from insert_chart import add_grunt_web_chart_shape
from restore_assets import read_assets, restore_assets


class RestoreAssetsTests(unittest.TestCase):
    def test_restores_exact_svg_and_safe_manifest_after_downloads_are_gone(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            # The token-looking text must remain only in the recovered SVG, never the manifest.
            svg = b'<svg xmlns="http://www.w3.org/2000/svg" data-grunt-src="secret-capability"><text>Chart + legend</text></svg>'
            tags = {"GRUNT-WEB-ASSETID": "12345678-1234-1234-1234-123456789abc",
                    "GRUNT-WEB-ASSETCOMMIT": "a" * 64, "GRUNT-WEB-LINK": "secret-capability",
                    "GRUNT-WEB-REVISION-POLICY": "pinned"}
            add_grunt_web_chart_shape(slide, slide.part, svg, b"fallback", 10, 20, 300, 200, "Chart", tags)
            source = root / "source.pptx"
            deck.save(source)
            manifest = restore_assets(source, root / "restored", root / "manifest.json")
            self.assertEqual(len(manifest["charts"]), 1)
            chart = manifest["charts"][0]
            self.assertEqual(chart["boundsEmu"], {"left": 10, "top": 20, "width": 300, "height": 200})
            self.assertEqual(chart["revisionPolicy"], "pinned")
            self.assertEqual((root / "restored" / chart["filename"]).read_bytes(), svg)
            self.assertNotIn("secret-capability", json.dumps(manifest))
            self.assertEqual(len(manifest["finalAssetIds"]), 1)
            self.assertEqual(read_assets(source)[0], manifest)

    def test_invalid_identity_cannot_be_used_as_an_output_filename(self):
        with tempfile.TemporaryDirectory() as directory:
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            add_grunt_web_chart_shape(slide, slide.part, b'<svg xmlns="http://www.w3.org/2000/svg"/>', b"fallback",
                                     10, 20, 30, 40, "Chart", {"GRUNT-WEB-ASSETID": "../escape"})
            path = Path(directory) / "invalid.pptx"
            deck.save(path)
            with self.assertRaisesRegex(ValueError, "Invalid Grunt identity"):
                restore_assets(path, Path(directory) / "restored")


if __name__ == "__main__":
    unittest.main()
