#!/usr/bin/env python3
"""
Tiny unit + functional test suite for this skill's scripts. No test framework
dependency (just `assert` + a hand-rolled runner) -- run directly with:

    python test_scripts.py

Covers:
  1. Unit tests of the shared parse/validate functions (parse_carrier_url,
     is_canonical_base32_secret, is_structurally_valid_token, parse_svg_metadata,
     redact_url) against fixture SVG bytes -- including adversarial fixtures
     (a CDATA-embedded lookalike comment, a comment outside the root, a comment
     nested inside a child element, malformed/duplicate query params, embedded
     credentials, a fragment) that must all be REJECTED. Always run, no external
     dependencies beyond the stdlib.
  2. Mocked-opener security tests of `download_svg` itself: HTTPS enforcement
     (no network call attempted for a non-HTTPS URL), redirect refusal, MIME
     checking, the byte cap enforced against ACTUAL bytes read (not just a
     claimed Content-Length), write identity (file on disk == bytes returned),
     and -- for every induced failure path -- asserting the bearer token from
     the test fixture URL appears in NEITHER captured stdout NOR stderr.
  3. An end-to-end functional test of insert_chart.py: builds a minimal pptx +
     a fixture carrier SVG in a temp dir, runs the real insertion code path,
     then asserts all three GRUNT-WEB-* tags are present with the right values
     AND that the embedded SVG media part is byte-identical to the fixture (the
     regression test for the P0 comment-stripping bug this version fixes), plus
     that a malformed --asset-id override is rejected. Skipped with a clear
     message if python-pptx/lxml aren't importable in this environment.
"""

import argparse
import base64
import contextlib
import importlib.util
import io
import json
import math
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent

_failures = []
_passed = 0


def check(condition: bool, description: str) -> None:
    global _passed
    if condition:
        _passed += 1
    else:
        _failures.append(description)


def load_module(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_insert_chart():
    """Import insert_chart.py. Rasterization is unrelated to what this test file
    exercises (tag construction, byte-preserving SVG embedding, origin trust),
    so when neither resvg-py nor a loadable CairoSVG is present, the module's
    svg_to_png is replaced with a stub and every other check still runs for real.
    Returns None (with a SKIPPED message) if python-pptx/lxml aren't importable --
    those ARE required."""
    try:
        from pptx import Presentation  # noqa: F401
    except ImportError as exc:
        print(f"  SKIPPED (python-pptx not installed: {exc})")
        return None

    module = load_module("insert_chart")
    try:
        module.svg_to_png(b'<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1"/>', 1)
    except ValueError as exc:
        module.svg_to_png = lambda svg_bytes, width=2500: b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
        print(f"  NOTE: no SVG rasterizer ({exc}); substituting a stub rasterizer for "
              f"this run. Tag writing and byte-preserving SVG embedding are still "
              f"exercised for real.")
    return module


def run_capturing(fn, *args, **kwargs):
    """Runs fn, capturing stdout/stderr and any SystemExit -- returns
    (return_value_or_None, stdout_text, stderr_text, exit_code_or_None)."""
    out, err = io.StringIO(), io.StringIO()
    exit_code = None
    result = None
    try:
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            result = fn(*args, **kwargs)
    except SystemExit as e:
        exit_code = e.code
    return result, out.getvalue(), err.getvalue(), exit_code


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

FIXTURE_ASSET_ID = "557db558-1234-4abc-9def-0123456789ab"
FIXTURE_COMMIT = "9a337151" + "0" * 56  # 64 hex chars

# A genuinely canonical Base32 secret (verified against this skill's own
# is_canonical_base32_secret below, not just assumed) -- base32(bytes(range(32)))
# with padding stripped. Deterministic and obviously not a real secret.
FIXTURE_SECRET = base64.b32encode(bytes(range(32))).decode("ascii").rstrip("=")
assert len(FIXTURE_SECRET) == 52
FIXTURE_TOKEN = ("a" * 32) + "." + FIXTURE_SECRET
FIXTURE_LINK = (
    f"https://assets.grunt.io/{FIXTURE_ASSET_ID}"
    f"?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"
)

# `&` in an XML attribute value must be escaped as `&amp;` -- a real server-rendered
# SVG does this (it's the only way the SVG itself is valid XML); ElementTree
# transparently un-escapes it back to a literal `&` on .get().
CARRIER_SVG = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="800pt" height="450pt"
     data-grunt-link-version="1"
     data-grunt-src="{FIXTURE_LINK.replace('&', '&amp;')}">
  <!-- a harmless in-root comment that must NOT break parsing -->
  <rect width="800" height="450" fill="#fff"/>
</svg>
""".encode("utf-8")

LEGACY_COMMENT_SVG = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 640 360">
  <!-- assetId:{FIXTURE_ASSET_ID},commitId:{FIXTURE_COMMIT} -->
  <rect width="640" height="360" fill="#fff"/>
</svg>
""".encode("utf-8")

NO_CARRIER_SVG = b"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50">
  <rect width="100" height="50" fill="#fff"/>
</svg>
"""

# Adversarial: the identity-shaped text sits inside a CDATA section, a nested
# child element, or the document epilog (after the root closes) -- none of
# these are a genuine direct-child comment of the root <svg>, and the fixed
# implementation must find NOTHING in any of them.
CDATA_LOOKALIKE_SVG = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50">
  <script><![CDATA[
    // <!-- assetId:{FIXTURE_ASSET_ID},commitId:{FIXTURE_COMMIT} -->
  ]]></script>
  <rect width="100" height="50" fill="#fff"/>
</svg>
""".encode("utf-8")

NESTED_COMMENT_SVG = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50">
  <g>
    <!-- assetId:{FIXTURE_ASSET_ID},commitId:{FIXTURE_COMMIT} -->
    <rect width="100" height="50" fill="#fff"/>
  </g>
</svg>
""".encode("utf-8")

EPILOG_COMMENT_SVG = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50">
  <rect width="100" height="50" fill="#fff"/>
</svg>
<!-- assetId:{FIXTURE_ASSET_ID},commitId:{FIXTURE_COMMIT} -->
""".encode("utf-8")

# A carrier whose origin is deliberately NOT in DEFAULT_ASSET_LINK_ORIGINS --
# the fixture for --trust-origin: this SVG's carrier must be invisible
# without the flag and recognized once that exact origin is trusted (the
# live-demo-discovered bug this version fixes: any non-production capability
# server's carrier used to be silently invisible with no override at all).
# Also carries the legacy identity comment, matching a real server-rendered
# SVG (both are always present together) -- this is what lets identity still
# resolve 2-of-3 (asset_id/asset_commit, no link) via that fallback when the
# carrier's origin isn't trusted, exactly as the bug report describes, rather
# than insert_chart.py finding no identity at all and refusing to proceed.
TRUSTED_ORIGIN = "https://localhost:7299"
TRUSTED_ORIGIN_LINK = (
    f"{TRUSTED_ORIGIN}/{FIXTURE_ASSET_ID}"
    f"?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"
)
TRUSTED_ORIGIN_SVG = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="800pt" height="450pt"
     data-grunt-link-version="1"
     data-grunt-src="{TRUSTED_ORIGIN_LINK.replace('&', '&amp;')}">
  <!-- assetId:{FIXTURE_ASSET_ID},commitId:{FIXTURE_COMMIT} -->
  <rect width="800" height="450" fill="#fff"/>
</svg>
""".encode("utf-8")


# ---------------------------------------------------------------------------
# 1. Unit tests
# ---------------------------------------------------------------------------

def test_download_chart_unit():
    dc = load_module("download_chart")

    # -- token shape validation ----------------------------------------------
    check(dc.is_canonical_base32_secret(FIXTURE_SECRET),
          "is_canonical_base32_secret: accepts the fixture's real canonical secret")
    check(not dc.is_canonical_base32_secret("B" * 52),
          "is_canonical_base32_secret: rejects a noncanonical (nonzero-padding-bit) secret")
    check(not dc.is_canonical_base32_secret("1" * 52),
          "is_canonical_base32_secret: rejects characters outside the Base32 alphabet")
    check(dc.is_structurally_valid_token(FIXTURE_TOKEN),
          "is_structurally_valid_token: accepts the well-formed fixture token")
    check(not dc.is_structurally_valid_token(FIXTURE_TOKEN + "x"),
          "is_structurally_valid_token: rejects a wrong-length secret")
    check(not dc.is_structurally_valid_token(("a" * 32) + "." + ("a" * 32) + "." + FIXTURE_SECRET),
          "is_structurally_valid_token: rejects more than one '.'")

    # -- parse_carrier_url: acceptance ----------------------------------------
    carrier = dc.parse_carrier_url(FIXTURE_LINK, dc.DEFAULT_ASSET_LINK_ORIGINS)
    check(carrier is not None, "parse_carrier_url: accepts a well-formed link")
    check(carrier and carrier["asset_id"] == FIXTURE_ASSET_ID.lower(),
          "parse_carrier_url: extracts the correct (lowercased) asset_id")
    check(carrier and carrier["asset_commit"] == FIXTURE_COMMIT.lower(),
          "parse_carrier_url: extracts the correct (lowercased) commit")
    check(carrier and carrier["link_id"] == "a" * 32,
          "parse_carrier_url: extracts the correct linkId (token's id segment)")

    # -- parse_carrier_url: adversarial rejections ----------------------------
    def carrier_url(**overrides):
        asset_id = overrides.get("asset_id", FIXTURE_ASSET_ID)
        commit = overrides.get("commit_qs", f"commit={FIXTURE_COMMIT}")
        token = overrides.get("token_qs", f"token={FIXTURE_TOKEN}")
        scheme = overrides.get("scheme", "https")
        host = overrides.get("host", "assets.grunt.io")
        userinfo = overrides.get("userinfo", "")
        fragment = overrides.get("fragment", "")
        return f"{scheme}://{userinfo}{host}/{asset_id}?{commit}&{token}{fragment}"

    check(dc.parse_carrier_url(carrier_url(scheme="http"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects non-HTTPS")
    check(dc.parse_carrier_url(carrier_url(host="evil.example.com"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects an origin not in the allowlist")
    check(dc.parse_carrier_url(carrier_url(userinfo="user:pass@"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects embedded credentials")
    check(dc.parse_carrier_url(carrier_url(fragment="#token=leak"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a fragment")
    check(dc.parse_carrier_url(carrier_url(asset_id="not-a-guid"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a malformed assetId path segment")
    check(dc.parse_carrier_url(carrier_url(commit_qs="commit=not-hex"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a non-hex commit")
    check(dc.parse_carrier_url(carrier_url(commit_qs=f"commit={FIXTURE_COMMIT}&commit={FIXTURE_COMMIT}"),
                                dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a duplicated commit query param")
    check(dc.parse_carrier_url(carrier_url(token_qs=f"token={FIXTURE_TOKEN}&token={FIXTURE_TOKEN}"),
                                dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a duplicated token query param")
    check(dc.parse_carrier_url(f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}",
                                dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a missing token")
    check(dc.parse_carrier_url(carrier_url(token_qs="token=" + ("a" * 32) + "." + ("B" * 52)),
                                dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a structurally-invalid (noncanonical) token")

    # -- parse_carrier_url: port 0 is a real, distinct port (falsy in Python, --
    # -- but not "no port") -- an explicit :0 must NOT reconstruct to the ------
    # -- SAME origin string as the bare default and sneak past the allowlist --
    check(dc.parse_carrier_url(carrier_url(host="api.grunt.io:0"), dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: an explicit :0 port is NOT treated as the bare default origin (live-verified bypass)")

    # -- parse_carrier_url: parse_qs(keep_blank_values=True) -- a blank -------
    # -- duplicate must be COUNTED (and so rejected as duplicate), never -------
    # -- silently dropped, leaving what looks like exactly one valid value ----
    check(dc.parse_carrier_url(
              carrier_url(commit_qs=f"commit={FIXTURE_COMMIT}&commit="),
              dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a blank duplicate commit query param, not silently dropped")
    check(dc.parse_carrier_url(
              carrier_url(token_qs=f"token={FIXTURE_TOKEN}&token="),
              dc.DEFAULT_ASSET_LINK_ORIGINS) is None,
          "parse_carrier_url: rejects a blank duplicate token query param, not silently dropped")

    # -- parse_svg_metadata: carrier attribute is primary --------------------
    meta = dc.parse_svg_metadata(CARRIER_SVG)
    check(meta["carrier_found"] is True, "parse_svg_metadata: finds the root carrier attribute")
    check(meta["asset_id"] == FIXTURE_ASSET_ID.lower(), "parse_svg_metadata: carrier asset_id correct")
    check(meta["asset_commit"] == FIXTURE_COMMIT.lower(), "parse_svg_metadata: carrier asset_commit correct")
    check(math.isclose(meta["width"], 800.0) and math.isclose(meta["height"], 450.0),
          "parse_svg_metadata: width/height parsed from attributes (pt suffix stripped)")

    # -- parse_svg_metadata: legacy comment fallback (genuine direct-child comment) --
    meta2 = dc.parse_svg_metadata(LEGACY_COMMENT_SVG)
    check(meta2["carrier_found"] is False,
          "parse_svg_metadata: legacy-comment-only SVG reports carrier_found=False")
    check(meta2["asset_id"] == FIXTURE_ASSET_ID, "parse_svg_metadata: legacy comment asset_id correct")
    check(meta2["asset_commit"] == FIXTURE_COMMIT, "parse_svg_metadata: legacy comment asset_commit correct")
    check(meta2["width"] == 480.0 and meta2["height"] == 270.0,
          "parse_svg_metadata: viewBox CSS pixels converted to points when width/height attrs absent")

    # -- parse_svg_metadata: no identity at all --------------------------------
    meta3 = dc.parse_svg_metadata(NO_CARRIER_SVG)
    check(meta3["asset_id"] is None and meta3["asset_commit"] is None,
          "parse_svg_metadata: no carrier/comment -> asset_id/commit are None")

    # -- parse_svg_metadata: adversarial lookalikes must all be REJECTED -------
    check(dc.parse_svg_metadata(CDATA_LOOKALIKE_SVG)["asset_id"] is None,
          "parse_svg_metadata: a CDATA-embedded lookalike is NOT treated as a comment")
    check(dc.parse_svg_metadata(NESTED_COMMENT_SVG)["asset_id"] is None,
          "parse_svg_metadata: a comment nested inside a child element is NOT 'inside the root'")
    check(dc.parse_svg_metadata(EPILOG_COMMENT_SVG)["asset_id"] is None,
          "parse_svg_metadata: a comment after the root element closes is NOT 'inside the root'")

    # -- final identity validation rejects malformed values regardless of source --
    forged_legacy = f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1 1">
  <!-- assetId:not-a-guid,commitId:not-hex -->
</svg>
""".encode("utf-8")
    meta4 = dc.parse_svg_metadata(forged_legacy)
    check(meta4["asset_id"] is None and meta4["asset_commit"] is None,
          "parse_svg_metadata: malformed legacy-comment values are rejected, not passed through")

    # -- redact_url ------------------------------------------------------------
    redacted = dc.redact_url(FIXTURE_LINK)
    check(FIXTURE_TOKEN not in redacted, "redact_url: strips the token out of the query string")
    check(redacted == f"https://assets.grunt.io/{FIXTURE_ASSET_ID}",
          "redact_url: keeps scheme+host+path only")
    redacted_creds = dc.redact_url(f"https://user:{FIXTURE_TOKEN}@api.grunt.io/{FIXTURE_ASSET_ID}")
    check(FIXTURE_TOKEN not in redacted_creds and "user" not in redacted_creds,
          "redact_url: strips embedded userinfo credentials from the netloc")

    # -- redact_url: port 0 is a real, distinct port -- must be PRESERVED, ----
    # -- never silently dropped the way bare `if parsed.port` would ------------
    redacted_port0 = dc.redact_url(f"https://assets.grunt.io:0/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}")
    check(redacted_port0 == f"https://assets.grunt.io:0/{FIXTURE_ASSET_ID}",
          "redact_url: preserves an explicit :0 port instead of silently dropping it")

    # -- redact_url: TOTAL -- an out-of-range port (whose `.port` access -------
    # -- itself raises ValueError) must fail to a fixed placeholder, NEVER -----
    # -- an uncaught traceback that echoes the offending port/token text -------
    # -- verbatim to stderr. This is exactly the shape of URL the new ----------
    # -- pre-fetch rejection path in download_svg redacts before printing. -----
    bad_port_url = f"https://assets.grunt.io:999999999/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"
    redacted_bad_port = dc.redact_url(bad_port_url)
    check(redacted_bad_port == "<unparseable URL>",
          "redact_url: an out-of-range port fails to the fixed placeholder, not a traceback")
    check(FIXTURE_TOKEN not in redacted_bad_port,
          "redact_url: the out-of-range-port fallback never echoes the token")

    # -- download_chart.py's OWN _validate_trust_origin: port 0 must be --------
    # -- PRESERVED (a real, distinct port), never silently dropped -- unlike ---
    # -- the shared normalization/rejection sweep in test_trust_origin_flag() -
    # -- (get_svg_info.py/insert_chart.py only), this is download_chart.py's --
    # -- own dedicated copy of the function -------------------------------------
    check(dc._validate_trust_origin("https://localhost:0") == "https://localhost:0",
          "download_chart._validate_trust_origin: an explicit :0 port is preserved, not silently dropped")

    # -- comments never break parsing (the P0 spike's core claim) --------------
    check(ET.fromstring(CARRIER_SVG).get("data-grunt-src") == FIXTURE_LINK,
          "ElementTree: in-root comment does not prevent attribute parsing")


def test_get_svg_info_unit():
    gi = load_module("get_svg_info")
    meta = gi.parse_svg_metadata(CARRIER_SVG)
    check(meta["carrier_found"] is True and meta["asset_id"] == FIXTURE_ASSET_ID.lower(),
          "get_svg_info.parse_svg_metadata: agrees with download_chart's own parser")
    check(gi.parse_svg_metadata(CDATA_LOOKALIKE_SVG)["asset_id"] is None,
          "get_svg_info.parse_svg_metadata: also rejects the CDATA lookalike")


def test_trust_origin_flag():
    """Regression coverage for the live-demo-discovered defect: against any
    non-default capability server, a carrier used to be silently invisible
    (carrier_found=False) with no way to override it. --trust-origin fixes
    that for get_svg_info.py and insert_chart.py.

    download_chart.py has its OWN --trust-origin coverage in
    test_download_svg_security() below (pre-fetch origin/path/commit/token
    validation, plus a CLI end-to-end case) -- it is NOT covered here.
    download_chart.py used to auto-trust whatever origin a download URL
    happened to point at (an SSRF/confused-deputy gap: any https origin
    could be fetched from, and would then be trusted for the post-download
    carrier check purely because it was the URL's own origin); it no longer
    does, so it needs its own dedicated coverage rather than being lumped in
    with get_svg_info.py/insert_chart.py's shared parse_svg_metadata-level
    tests below."""
    dc = load_module("download_chart")
    gi = load_module("get_svg_info")
    ic = load_insert_chart()

    mods = [("get_svg_info", gi)]
    if ic is not None:
        mods.append(("insert_chart", ic))

    # -- hostile port: parse_carrier_url must fail closed (return None), --
    # -- never crash with a raw ValueError from accessing `.port` ----------
    # (present in ALL THREE modules -- download_chart.py, get_svg_info.py,
    # insert_chart.py -- since each has its own copy of parse_carrier_url).
    hostile_port_url = (
        f"https://assets.grunt.io:999999999/{FIXTURE_ASSET_ID}"
        f"?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"
    )
    for name, mod in [("download_chart", dc)] + mods:
        try:
            result = mod.parse_carrier_url(hostile_port_url, mod.DEFAULT_ASSET_LINK_ORIGINS)
            check(result is None,
                  f"{name}.parse_carrier_url: an out-of-range port fails closed (returns None), doesn't crash")
        except Exception as exc:
            check(False, f"{name}.parse_carrier_url: an out-of-range port must not raise ({type(exc).__name__}: {exc})")

    # -- explicit :0 port must NOT be treated as "no port" (falsy != absent) --
    # -- an origin with :0 must NOT reconstruct to the SAME origin string as --
    # -- the bare default and sneak past the allowlist (present in ALL THREE -
    # -- modules, same reasoning as the hostile-port check above -- live- -----
    # -- verified bypass: https://assets.grunt.io:0/... used to pass the gate ----
    # -- as an impostor of the real https://assets.grunt.io default) -------------
    port0_url = (
        f"https://assets.grunt.io:0/{FIXTURE_ASSET_ID}"
        f"?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"
    )
    for name, mod in [("download_chart", dc)] + mods:
        check(mod.parse_carrier_url(port0_url, mod.DEFAULT_ASSET_LINK_ORIGINS) is None,
              f"{name}.parse_carrier_url: an explicit :0 port is NOT treated as the bare default origin")

    # -- a blank duplicate commit/token must be COUNTED (and so rejected as ---
    # -- a duplicate), never silently dropped by parse_qs's default -----------
    # -- keep_blank_values=False behavior (present in ALL THREE modules) ------
    blank_commit_url = (
        f"https://assets.grunt.io/{FIXTURE_ASSET_ID}"
        f"?commit={FIXTURE_COMMIT}&commit=&token={FIXTURE_TOKEN}"
    )
    blank_token_url = (
        f"https://assets.grunt.io/{FIXTURE_ASSET_ID}"
        f"?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}&token="
    )
    for name, mod in [("download_chart", dc)] + mods:
        check(mod.parse_carrier_url(blank_commit_url, mod.DEFAULT_ASSET_LINK_ORIGINS) is None,
              f"{name}.parse_carrier_url: rejects a blank duplicate commit query param, not silently dropped")
        check(mod.parse_carrier_url(blank_token_url, mod.DEFAULT_ASSET_LINK_ORIGINS) is None,
              f"{name}.parse_carrier_url: rejects a blank duplicate token query param, not silently dropped")

    for name, mod in mods:
        # -- _validate_trust_origin: normalization ---------------------------
        check(mod._validate_trust_origin("https://localhost:7299") == "https://localhost:7299",
              f"{name}._validate_trust_origin: a bare https origin is returned unchanged")
        check(mod._validate_trust_origin("https://localhost:7299/") == "https://localhost:7299",
              f"{name}._validate_trust_origin: an allowed trailing slash is stripped")
        check(mod._validate_trust_origin("https://LOCALHOST:7299") == "https://localhost:7299",
              f"{name}._validate_trust_origin: the host is lowercased")
        check(mod._validate_trust_origin("https://assets.grunt.io") == "https://assets.grunt.io",
              f"{name}._validate_trust_origin: an origin with no port is returned unchanged")
        check(mod._validate_trust_origin("https://localhost:0") == "https://localhost:0",
              f"{name}._validate_trust_origin: an explicit :0 port is preserved, not silently dropped")

        # -- _validate_trust_origin: rejections (never a silent no-op) -------
        for bad in [
            "http://localhost:7299",            # non-HTTPS (no exception for localhost)
            "https://localhost:7299/path",       # path beyond a bare trailing slash
            "https://user:pw@localhost:7299",    # embedded credentials
            "https://localhost:7299?x=1",        # query string
            "https://localhost:7299#frag",       # fragment
            "https:///only-a-path",       # no host
            "not-a-url-at-all",                  # no scheme/host at all
            "https://localhost:7299 ",           # trailing whitespace
            "https://local host:7299",           # embedded whitespace
            "https://localhost:7299\t",          # embedded control char (tab)
            "https://localhost:7299\n",          # embedded control char (newline)
            f"https://localhost:999999999",      # out-of-range port
        ]:
            try:
                mod._validate_trust_origin(bad)
                check(False, f"{name}._validate_trust_origin: should reject {bad!r}")
            except argparse.ArgumentTypeError:
                check(True, f"{name}._validate_trust_origin: rejects {bad!r}")

        # -- _validate_trust_origin: NEVER echoes the candidate value back in --
        # -- its error message -- a caller could mistakenly pass an entire ------
        # -- token-bearing capability URL here instead of a bare origin --------
        try:
            mod._validate_trust_origin(FIXTURE_LINK)
            check(False, f"{name}._validate_trust_origin: should reject a full token-bearing URL")
        except argparse.ArgumentTypeError as exc:
            check(FIXTURE_TOKEN not in str(exc),
                  f"{name}._validate_trust_origin: a mistakenly-passed token-bearing URL's "
                  f"token never appears in the ArgumentTypeError message")

        # -- parse_svg_metadata: carrier from a non-default origin ------------
        without = mod.parse_svg_metadata(TRUSTED_ORIGIN_SVG)
        check(without["carrier_found"] is False if name == "get_svg_info" else without["link"] is None,
              f"{name}.parse_svg_metadata: a non-default-origin carrier is NOT recognized without --trust-origin")

        with_trust = mod.parse_svg_metadata(TRUSTED_ORIGIN_SVG, extra_trusted_origins=(TRUSTED_ORIGIN,))
        if name == "get_svg_info":
            check(with_trust["carrier_found"] is True and with_trust["asset_id"] == FIXTURE_ASSET_ID.lower(),
                  f"{name}.parse_svg_metadata: the carrier IS recognized once its origin is trusted via extra_trusted_origins")
        else:
            check(with_trust["link"] == TRUSTED_ORIGIN_LINK,
                  f"{name}.parse_svg_metadata: the link IS emitted once its origin is trusted via extra_trusted_origins")

        # -- the default origin is unaffected: still works with no extra origins --
        default_still_works = mod.parse_svg_metadata(CARRIER_SVG)
        check(default_still_works["asset_id"] == FIXTURE_ASSET_ID.lower(),
              f"{name}.parse_svg_metadata: the built-in default origin still works with no extra_trusted_origins")

    # -- get_svg_info.py CLI end-to-end: without / with / invalid --trust-origin --
    with tempfile.TemporaryDirectory() as tmp:
        svg_path = Path(tmp) / "trusted.svg"
        svg_path.write_bytes(TRUSTED_ORIGIN_SVG)

        old_argv = sys.argv
        try:
            sys.argv = ["get_svg_info.py", str(svg_path)]
            _, out, err, code = run_capturing(gi.main)
        finally:
            sys.argv = old_argv
        check(code is None, "get_svg_info.py CLI: runs cleanly with no --trust-origin")
        result = json.loads(out)
        check(result["carrier_found"] is False,
              "get_svg_info.py CLI: without --trust-origin, a non-default-origin carrier is not found")

        old_argv = sys.argv
        try:
            sys.argv = ["get_svg_info.py", str(svg_path), "--trust-origin", TRUSTED_ORIGIN]
            _, out, err, code = run_capturing(gi.main)
        finally:
            sys.argv = old_argv
        check(code is None, "get_svg_info.py CLI: runs cleanly with --trust-origin")
        result = json.loads(out)
        check(result["carrier_found"] is True and result["asset_id"] == FIXTURE_ASSET_ID.lower(),
              "get_svg_info.py CLI: --trust-origin extends the allowlist and the carrier is now recognized")

        old_argv = sys.argv
        try:
            sys.argv = ["get_svg_info.py", str(svg_path), "--trust-origin", "http://localhost:7299"]
            _, out, err, code = run_capturing(gi.main)
        finally:
            sys.argv = old_argv
        check(code not in (None, 0),
              "get_svg_info.py CLI: an invalid --trust-origin value is rejected (nonzero exit), never silently ignored")

        # -- CLI-level redaction: a mistakenly-pasted token-bearing URL as ---
        # -- --trust-origin's value must never echo the token to stdout OR --
        # -- stderr, whether from our own message or argparse's own fallback -
        old_argv = sys.argv
        try:
            sys.argv = ["get_svg_info.py", str(svg_path), "--trust-origin", FIXTURE_LINK]
            _, out, err, code = run_capturing(gi.main)
        finally:
            sys.argv = old_argv
        check(code not in (None, 0),
              "get_svg_info.py CLI: a token-bearing invalid --trust-origin value is rejected")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "get_svg_info.py CLI: the token from a mistakenly-pasted full URL never echoes "
              "to stdout or stderr")


# ---------------------------------------------------------------------------
# 2. Mocked-opener security tests of download_svg
# ---------------------------------------------------------------------------

class _FakeHeaders:
    def __init__(self, mapping):
        self._mapping = {k.lower(): v for k, v in mapping.items()}

    def get(self, key, default=None):
        return self._mapping.get(key.lower(), default)


class _FakeResponse:
    def __init__(self, body: bytes, headers: dict):
        self._body = body
        self.headers = _FakeHeaders(headers)

    def read(self, n=-1):
        if n is None or n < 0:
            data, self._body = self._body, b""
            return data
        data, self._body = self._body[:n], self._body[n:]
        return data

    def __enter__(self):
        return self

    def __exit__(self, *exc_info):
        return False


class _FakeOpener:
    def __init__(self, response=None, error=None):
        self._response = response
        self._error = error
        self.opened = False

    def open(self, req, timeout=None):
        self.opened = True
        if self._error is not None:
            raise self._error
        return self._response


class _ManifestOpener:
    def __init__(self, svg_response, source):
        self.svg_response = svg_response
        self.source = source

    def open(self, req, timeout=None):
        if req.get_header("Accept") == "image/svg+xml":
            return self.svg_response
        body = json.dumps({"schemaVersion": 1, "assetId": FIXTURE_ASSET_ID.lower(),
            "resolvedCommit": FIXTURE_COMMIT, "links": {
                "self": {"href": self.source, "type": "application/json", "schemaVersion": 1},
                "svg": {"href": self.source.split("/", 3)[0] + "//" + self.source.split("/", 3)[2] + "/assets/svg?opaque=delivery", "type": "image/svg+xml"}
            }}).encode()
        return _FakeResponse(body, {"Content-Type": "application/json"})


def test_download_svg_security():
    dc = load_module("download_chart")

    test_url = f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"

    with tempfile.TemporaryDirectory() as tmp:
        out_path = str(Path(tmp) / "out.svg")

        # -- zero-network guard: records every invocation into `calls` BEFORE ---
        # -- raising, so "no request attempted" can be asserted on `len(calls)` -
        # -- AFTER run_capturing catches the SystemExit -- independent of ------
        # -- whatever exception handling download_svg does internally. A bare --
        # -- AssertionError-raising stub is NOT sufficient on its own: --------
        # -- download_svg's own broad `except Exception` catches it and turns --
        # -- it into the exact same fail()/exit-1 outcome the rejection is -----
        # -- already expected to produce, so checking `code == 1` alone can't --
        # -- distinguish "rejected before any request" from "the guard fired ---
        # -- because a request WAS attempted and its exception got swallowed" --
        original_build_opener = dc.build_opener

        def _make_guard():
            calls = []

            def _opener_that_must_not_be_called(*a, **kw):
                calls.append((a, kw))
                raise AssertionError("build_opener must not be called for a rejected URL")

            return calls, _opener_that_must_not_be_called

        # -- HTTPS enforcement: no network call attempted for non-HTTPS --------
        calls, guard = _make_guard()
        dc.build_opener = guard
        try:
            _, out, err, code = run_capturing(
                dc.download_svg,
                f"http://api.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}",
                out_path, dc.DEFAULT_ASSET_LINK_ORIGINS,
            )
        finally:
            dc.build_opener = original_build_opener
        check(code == 1, "download_svg: non-HTTPS URL exits nonzero")
        check(len(calls) == 0, "download_svg: non-HTTPS rejection makes NO network attempt")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "download_svg: non-HTTPS rejection never echoes the token (stdout+stderr)")

        # -- credentials / fragment rejected before any request -----------------
        _, out, err, code = run_capturing(
            dc.download_svg,
            f"https://user:pw@api.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}",
            out_path, dc.DEFAULT_ASSET_LINK_ORIGINS,
        )
        check(code == 1 and "pw" not in out and "pw" not in err,
              "download_svg: embedded credentials rejected, never echoed")

        _, out, err, code = run_capturing(
            dc.download_svg, test_url + "#leak", out_path, dc.DEFAULT_ASSET_LINK_ORIGINS
        )
        check(code == 1, "download_svg: URL fragment rejected")

        # -- pre-fetch full-carrier validation: wrong origin, wrong path shape, --
        # -- missing/malformed commit/token, an explicit :0-port impostor of ----
        # -- the default origin, and an out-of-range port (which must fail to --
        # -- a clean fixed-category error via redact_url's own guard, never a --
        # -- traceback that echoes the port/token text) are ALL rejected -------
        # -- BEFORE any request is made (the SSRF-closing fix -- previously ----
        # -- only scheme/host/credentials/fragment were checked before ANY -----
        # -- request went out) --------------------------------------------------
        pre_fetch_rejections = {
            "an origin not in the allowlist (no --trust-origin)":
                f"https://evil.example.com/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}",
            "a path that isn't /{uuid}":
                f"https://assets.grunt.io/not-assets/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}",
            "a malformed (non-hex) commit":
                f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit=not-hex&token={FIXTURE_TOKEN}",
            "a blank duplicate commit (keep_blank_values)":
                f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&commit=&token={FIXTURE_TOKEN}",
            "a missing token":
                f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}",
            "a structurally malformed token":
                f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token=not-a-real-token",
            "a blank duplicate token (keep_blank_values)":
                f"https://assets.grunt.io/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}&token=",
            "an explicit :0 port impostor of the default origin":
                f"https://assets.grunt.io:0/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}",
            "an out-of-range port (redact_url's own .port access must not raise)":
                f"https://assets.grunt.io:999999999/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}",
        }
        for description, bad_url in pre_fetch_rejections.items():
            calls, guard = _make_guard()
            dc.build_opener = guard
            try:
                _, out, err, code = run_capturing(
                    dc.download_svg, bad_url, out_path, dc.DEFAULT_ASSET_LINK_ORIGINS
                )
            finally:
                dc.build_opener = original_build_opener
            check(code == 1, f"download_svg: pre-fetch validation rejects {description} (nonzero exit)")
            check(len(calls) == 0, f"download_svg: pre-fetch rejection of {description} makes NO network attempt")
            check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
                  f"download_svg: pre-fetch rejection of {description} never echoes the token")

        # -- --trust-origin: a non-default origin is rejected without it (no -----
        # -- request attempted), and admitted (request IS made) with it ---------
        trusted_url = f"{TRUSTED_ORIGIN}/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"
        calls, guard = _make_guard()
        dc.build_opener = guard
        try:
            _, out, err, code = run_capturing(
                dc.download_svg, trusted_url, out_path, dc.DEFAULT_ASSET_LINK_ORIGINS
            )
        finally:
            dc.build_opener = original_build_opener
        check(code == 1, "download_svg: a non-default origin is rejected without --trust-origin (nonzero exit)")
        check(len(calls) == 0, "download_svg: a non-default origin without --trust-origin makes NO network attempt")

        trusted_body = TRUSTED_ORIGIN_SVG
        fake = _FakeResponse(trusted_body, {"Content-Type": "image/svg+xml", "Content-Length": str(len(trusted_body))})
        dc.build_opener = lambda *a, **kw: _ManifestOpener(fake, trusted_url)
        try:
            returned, out, err, code = run_capturing(
                dc.download_svg, trusted_url, out_path,
                dc.DEFAULT_ASSET_LINK_ORIGINS + (TRUSTED_ORIGIN,),
            )
        finally:
            dc.build_opener = original_build_opener
        check(code is None, "download_svg: --trust-origin admits a matching non-default origin (request is attempted)")
        check(returned == trusted_body, "download_svg: the trusted non-default-origin download succeeds")

        # -- redirect refusal (direct unit test of the handler) -----------------
        handler = dc._NoRedirectHandler()
        try:
            handler.redirect_request(
                req=type("Req", (), {"full_url": test_url})(),
                fp=None, code=302, msg="Found", headers={}, newurl="https://evil.example.com/"
            )
            check(False, "redirect_request: should have raised")
        except Exception as e:
            check(isinstance(e, dc.HTTPError), "redirect_request: raises HTTPError to refuse the redirect")

        # -- MIME check ----------------------------------------------------------
        fake = _FakeResponse(b"<html>not an svg</html>", {"Content-Type": "text/html"})
        dc.build_opener = lambda *a, **kw: _FakeOpener(response=fake)
        try:
            _, out, err, code = run_capturing(dc.download_svg, test_url, out_path, dc.DEFAULT_ASSET_LINK_ORIGINS)
        finally:
            dc.build_opener = original_build_opener
        check(code == 1, "download_svg: wrong Content-Type is rejected")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "download_svg: MIME rejection never echoes the token")

        # -- actual-byte cap (not just a claimed Content-Length) -----------------
        original_cap = dc.MAX_RESPONSE_BYTES
        dc.MAX_RESPONSE_BYTES = 100
        try:
            oversized_body = b"<svg>" + (b"x" * 500) + b"</svg>"
            fake = _FakeResponse(oversized_body, {"Content-Type": "application/json"})  # no Content-Length header at all
            dc.build_opener = lambda *a, **kw: _FakeOpener(response=fake)
            _, out, err, code = run_capturing(dc.download_svg, test_url, out_path, dc.DEFAULT_ASSET_LINK_ORIGINS)
        finally:
            dc.MAX_RESPONSE_BYTES = original_cap
            dc.build_opener = original_build_opener
        check(code == 1, "download_svg: oversized ACTUAL body is rejected even with no Content-Length claim")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "download_svg: size-cap rejection never echoes the token")

        # -- HTTPError / URLError paths never echo raw reason text ---------------
        hostile_reason = f"redirect to {test_url} please"
        fake_error = dc.HTTPError(test_url, 500, hostile_reason, {}, None)
        dc.build_opener = lambda *a, **kw: _FakeOpener(error=fake_error)
        try:
            _, out, err, code = run_capturing(dc.download_svg, test_url, out_path, dc.DEFAULT_ASSET_LINK_ORIGINS)
        finally:
            dc.build_opener = original_build_opener
        check(code == 1, "download_svg: HTTPError path fails closed")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err
              and hostile_reason not in out and hostile_reason not in err,
              "download_svg: HTTPError's raw reason text is never echoed (only a fixed category + status code)")

        # -- write identity + successful path ------------------------------------
        good_body = CARRIER_SVG
        fake = _FakeResponse(good_body, {"Content-Type": "image/svg+xml", "Content-Length": str(len(good_body))})
        dc.build_opener = lambda *a, **kw: _ManifestOpener(fake, test_url)
        try:
            returned, out, err, code = run_capturing(dc.download_svg, test_url, out_path, dc.DEFAULT_ASSET_LINK_ORIGINS)
        finally:
            dc.build_opener = original_build_opener
        check(code is None, "download_svg: a valid response does not exit early")
        check(returned == good_body, "download_svg: returns the exact bytes received")
        check(Path(out_path).read_bytes() == good_body,
              "download_svg: the file written to disk is byte-identical to the response")


def test_download_chart_cli_trust_origin():
    """CLI end-to-end coverage for download_chart.py's --trust-origin, mirroring
    get_svg_info.py's CLI coverage in test_trust_origin_flag() above but for the
    PRE-fetch check (get_svg_info.py has no network request to gate at all).
    Runs the real main() (stdin-fed URL, mocked opener -- no real network I/O):
    (1) a non-default-origin URL is refused with NO request attempted when
    --trust-origin is omitted; (2) the SAME URL succeeds -- both the fetch and
    the post-download carrier recognition -- once --trust-origin names its
    origin; (3) an invalid --trust-origin value is rejected outright (nonzero
    exit, no token echoed, no request attempted either).

    The "no request attempted" assertions use a call-recording guard (see
    _make_guard in test_download_svg_security's own comment for why a bare
    AssertionError-raising stub is not sufficient on its own -- main() calls
    download_svg(), whose broad `except Exception` would swallow that
    AssertionError into the exact same fail()/exit-1 outcome the rejection
    already produces, so `code == 1` alone can't prove no request was made)."""
    dc = load_module("download_chart")
    original_build_opener = dc.build_opener
    trusted_url = f"{TRUSTED_ORIGIN}/{FIXTURE_ASSET_ID}?commit={FIXTURE_COMMIT}&token={FIXTURE_TOKEN}"

    def _make_guard():
        calls = []

        def _must_not_open(*a, **kw):
            calls.append((a, kw))
            raise AssertionError("build_opener must not be called for a rejected URL")

        return calls, _must_not_open

    with tempfile.TemporaryDirectory() as tmp:
        out_path = str(Path(tmp) / "cli_out.svg")

        # -- (1) without --trust-origin: refused, no request attempted ----------
        calls, guard = _make_guard()
        dc.build_opener = guard
        old_argv, old_stdin = sys.argv, sys.stdin
        try:
            sys.argv = ["download_chart.py", "--output", out_path]
            sys.stdin = io.StringIO(trusted_url + "\n")
            _, out, err, code = run_capturing(dc.main)
        finally:
            sys.argv, sys.stdin = old_argv, old_stdin
            dc.build_opener = original_build_opener
        check(code == 1, "download_chart.py CLI: a non-default-origin URL is refused without --trust-origin")
        check(len(calls) == 0, "download_chart.py CLI: the refusal above makes NO network attempt")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "download_chart.py CLI: the rejection never echoes the token")

        # -- (2) with --trust-origin: the request IS made and succeeds, and the --
        # -- post-download carrier is ALSO recognized via the same flag ---------
        good_body = TRUSTED_ORIGIN_SVG
        fake = _FakeResponse(good_body, {"Content-Type": "image/svg+xml", "Content-Length": str(len(good_body))})
        dc.build_opener = lambda *a, **kw: _ManifestOpener(fake, trusted_url)
        old_argv, old_stdin = sys.argv, sys.stdin
        try:
            sys.argv = ["download_chart.py", "--output", out_path, "--trust-origin", TRUSTED_ORIGIN]
            sys.stdin = io.StringIO(trusted_url + "\n")
            _, out, err, code = run_capturing(dc.main)
        finally:
            sys.argv, sys.stdin = old_argv, old_stdin
            dc.build_opener = original_build_opener
        check(code is None, "download_chart.py CLI: --trust-origin admits the matching non-default origin")
        result = json.loads(out)
        check(result.get("carrier_found") is True and result.get("asset_id") == FIXTURE_ASSET_ID.lower(),
              "download_chart.py CLI: the post-download carrier is also recognized via the same --trust-origin")

        # -- (3) an invalid --trust-origin value is rejected outright, no request -
        calls, guard = _make_guard()
        dc.build_opener = guard
        old_argv, old_stdin = sys.argv, sys.stdin
        try:
            sys.argv = ["download_chart.py", "--output", out_path, "--trust-origin", "http://localhost:7299"]
            sys.stdin = io.StringIO(trusted_url + "\n")
            _, out, err, code = run_capturing(dc.main)
        finally:
            sys.argv, sys.stdin = old_argv, old_stdin
            dc.build_opener = original_build_opener
        check(code not in (None, 0), "download_chart.py CLI: an invalid --trust-origin value is rejected")
        check(len(calls) == 0, "download_chart.py CLI: an invalid --trust-origin value makes NO network attempt")

        # -- CLI-level redaction: a mistakenly-pasted token-bearing URL as -------
        # -- --trust-origin's value must never echo the token to stdout OR ------
        # -- stderr, whether from our own message or argparse's own fallback ----
        calls, guard = _make_guard()
        dc.build_opener = guard
        old_argv, old_stdin = sys.argv, sys.stdin
        try:
            sys.argv = ["download_chart.py", "--output", out_path, "--trust-origin", FIXTURE_LINK]
            sys.stdin = io.StringIO(trusted_url + "\n")
            _, out, err, code = run_capturing(dc.main)
        finally:
            sys.argv, sys.stdin = old_argv, old_stdin
            dc.build_opener = original_build_opener
        check(code not in (None, 0), "download_chart.py CLI: a token-bearing invalid --trust-origin value is rejected")
        check(len(calls) == 0,
              "download_chart.py CLI: a token-bearing invalid --trust-origin value makes NO network attempt")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "download_chart.py CLI: the token from a mistakenly-pasted full URL never echoes "
              "to stdout or stderr")

        # -- CLI-level redaction: a token-bearing URL passed POSITIONALLY --------
        # -- (instead of via stdin/--url-file, the only supported channels) -----
        # -- must not be echoed by argparse's own "unrecognized arguments" ------
        # -- error path -- verified live to otherwise leak the token to stderr --
        calls, guard = _make_guard()
        dc.build_opener = guard
        old_argv, old_stdin = sys.argv, sys.stdin
        try:
            sys.argv = ["download_chart.py", "--output", out_path, FIXTURE_LINK]
            sys.stdin = io.StringIO(trusted_url + "\n")
            _, out, err, code = run_capturing(dc.main)
        finally:
            sys.argv, sys.stdin = old_argv, old_stdin
            dc.build_opener = original_build_opener
        check(code == 2, "download_chart.py CLI: a positional URL argument is rejected by argparse (exit 2)")
        check(len(calls) == 0, "download_chart.py CLI: a rejected positional argument makes NO network attempt")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "download_chart.py CLI: a token-bearing positional argument is never echoed by "
              "argparse's own 'unrecognized arguments' error (stdout or stderr)")


# ---------------------------------------------------------------------------
# 3. Functional end-to-end test of insert_chart.py
# ---------------------------------------------------------------------------

def test_insert_chart_functional():
    ic = load_insert_chart()
    if ic is None:
        return
    from pptx import Presentation

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        svg_path = tmp_path / "chart.svg"
        svg_path.write_bytes(CARRIER_SVG)

        pptx_in = tmp_path / "in.pptx"
        pptx_out = tmp_path / "out.pptx"

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        prs.save(pptx_in)

        # -- malformed --asset-id override is rejected, not silently accepted ---
        old_argv = sys.argv
        try:
            sys.argv = [
                "insert_chart.py",
                "--svg", str(svg_path), "--pptx", str(pptx_in), "--output", str(pptx_out),
                "--asset-id", "not-a-guid",
            ]
            _, out, err, code = run_capturing(ic.main)
        finally:
            sys.argv = old_argv
        check(code not in (None, 0), "insert_chart.py: a malformed --asset-id override is rejected")

        # -- real insertion run ---------------------------------------------------
        old_argv = sys.argv
        try:
            sys.argv = [
                "insert_chart.py",
                "--svg", str(svg_path),
                "--pptx", str(pptx_in),
                "--output", str(pptx_out),
                "--slide", "0",
                "--left", "1.0", "--top", "1.0", "--width", "8.0", "--height", "4.5",
            ]
            ic.main()
        finally:
            sys.argv = old_argv

        check(pptx_out.exists(), "insert_chart.py: writes the output pptx")

        with zipfile.ZipFile(pptx_out) as z:
            names = z.namelist()

            # -- byte-preservation regression test (the P0 bug this version fixes) --
            svg_media = [n for n in names if n.startswith("ppt/media/") and n.endswith(".svg")]
            check(len(svg_media) == 1, "insert_chart.py: exactly one SVG media part written")
            if svg_media:
                embedded = z.read(svg_media[0])
                check(embedded == CARRIER_SVG,
                      "insert_chart.py: embedded SVG bytes are BYTE-IDENTICAL to the "
                      "original file (regression test for the comment-stripping bug)")

            # -- PNG fallback present --------------------------------------------
            png_media = [n for n in names if n.startswith("ppt/media/") and n.endswith(".png")]
            check(len(png_media) == 1, "insert_chart.py: exactly one PNG fallback media part written")

            # -- all three tags present with correct values ----------------------
            ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            tag_files = [n for n in names if "tags/tag" in n and n.endswith(".xml")]
            found_tags = {}
            for tf in tag_files:
                for t in ET.fromstring(z.read(tf)).findall(".//p:tag", ns):
                    found_tags[t.get("name")] = t.get("val")

            check(found_tags.get("GRUNT-WEB-REVISION-POLICY") == "pinned",
                  "insert_chart.py: new placements default to the saved version")
            check(found_tags.get("GRUNT-WEB-ASSETID") == FIXTURE_ASSET_ID.lower(),
                  "insert_chart.py: GRUNT-WEB-ASSETID tag correct")
            check(found_tags.get("GRUNT-WEB-ASSETCOMMIT") == FIXTURE_COMMIT.lower(),
                  "insert_chart.py: GRUNT-WEB-ASSETCOMMIT tag correct")
            check(found_tags.get("GRUNT-WEB-LINK") == FIXTURE_LINK,
                  "insert_chart.py: GRUNT-WEB-LINK tag correct (byte-identical to data-grunt-src)")


        old_argv = sys.argv
        try:
            sys.argv = ["insert_chart.py", "--svg", str(svg_path), "--pptx", str(pptx_in),
                        "--output", str(pptx_out), "--slide", "0", "--left", "1", "--top", "1",
                        "--width", "8", "--height", "4.5", "--revision-policy", "latest"]
            run_capturing(ic.main)
        finally:
            sys.argv = old_argv
        with zipfile.ZipFile(pptx_out) as z:
            policies = [tag.get("val") for name in z.namelist() if "tags/tag" in name and name.endswith(".xml")
                        for tag in ET.fromstring(z.read(name)).findall(".//p:tag", ns)
                        if tag.get("name") == "GRUNT-WEB-REVISION-POLICY"]
            check(policies == ["latest"], "insert_chart.py: explicit latest opts into first-import refresh")


# ---------------------------------------------------------------------------
# 5. Functional test: insert_chart.py + --trust-origin
# ---------------------------------------------------------------------------

def test_insert_chart_trust_origin():
    """Regression test for the live-demo-discovered defect: against a chart
    carried from a non-default capability server, insert_chart.py used to
    NEVER write GRUNT-WEB-LINK (no override existed) -- the freshly inserted
    shape then failed the SKILL.md validation snippet with
    missing_or_empty={'GRUNT-WEB-LINK'} on every single insertion. This
    confirms --trust-origin fixes that end-to-end: absent without the flag,
    present (and byte-identical to the carrier) with it."""
    ic = load_insert_chart()
    if ic is None:
        return
    from pptx import Presentation

    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}

    def _tags_in(pptx_path):
        found = {}
        with zipfile.ZipFile(pptx_path) as z:
            for tf in [n for n in z.namelist() if "tags/tag" in n and n.endswith(".xml")]:
                for t in ET.fromstring(z.read(tf)).findall(".//p:tag", ns):
                    found[t.get("name")] = t.get("val")
        return found

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        svg_path = tmp_path / "chart.svg"
        svg_path.write_bytes(TRUSTED_ORIGIN_SVG)

        pptx_in = tmp_path / "in.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        prs.save(pptx_in)

        # -- WITHOUT --trust-origin: GRUNT-WEB-LINK must be absent (2 of 3) ------
        pptx_out_untrusted = tmp_path / "out_untrusted.pptx"
        old_argv = sys.argv
        try:
            sys.argv = [
                "insert_chart.py",
                "--svg", str(svg_path), "--pptx", str(pptx_in), "--output", str(pptx_out_untrusted),
                "--slide", "0", "--left", "1.0", "--top", "1.0", "--width", "8.0", "--height", "4.5",
            ]
            _, out, err, code = run_capturing(ic.main)
        finally:
            sys.argv = old_argv
        check(code is None, "insert_chart.py: runs cleanly on an untrusted-origin SVG with no --trust-origin")

        tags_untrusted = _tags_in(pptx_out_untrusted)
        check("GRUNT-WEB-LINK" not in tags_untrusted,
              "insert_chart.py: GRUNT-WEB-LINK is NOT written for an untrusted origin without --trust-origin "
              "(this is the live-demo-discovered bug -- validation would fail on this exact shape)")
        check(tags_untrusted.get("GRUNT-WEB-ASSETID") == FIXTURE_ASSET_ID.lower()
              and tags_untrusted.get("GRUNT-WEB-ASSETCOMMIT") == FIXTURE_COMMIT.lower(),
              "insert_chart.py: GRUNT-WEB-ASSETID/ASSETCOMMIT are still written (identity resolves 2-of-3)")

        # -- WITH --trust-origin: GRUNT-WEB-LINK IS written, byte-identical ------
        pptx_out_trusted = tmp_path / "out_trusted.pptx"
        old_argv = sys.argv
        try:
            sys.argv = [
                "insert_chart.py",
                "--svg", str(svg_path), "--pptx", str(pptx_in), "--output", str(pptx_out_trusted),
                "--slide", "0", "--left", "1.0", "--top", "1.0", "--width", "8.0", "--height", "4.5",
                "--trust-origin", TRUSTED_ORIGIN,
            ]
            _, out, err, code = run_capturing(ic.main)
        finally:
            sys.argv = old_argv
        check(code is None, "insert_chart.py: runs cleanly with --trust-origin")

        tags_trusted = _tags_in(pptx_out_trusted)
        check(tags_trusted.get("GRUNT-WEB-LINK") == TRUSTED_ORIGIN_LINK,
              "insert_chart.py: GRUNT-WEB-LINK IS written (byte-identical to data-grunt-src) once the "
              "carrier's origin is trusted via --trust-origin")
        check(tags_trusted.get("GRUNT-WEB-ASSETID") == FIXTURE_ASSET_ID.lower()
              and tags_trusted.get("GRUNT-WEB-ASSETCOMMIT") == FIXTURE_COMMIT.lower(),
              "insert_chart.py: GRUNT-WEB-ASSETID/ASSETCOMMIT still correct with --trust-origin")

        # -- an invalid --trust-origin value is rejected, never silently ignored --
        old_argv = sys.argv
        try:
            sys.argv = [
                "insert_chart.py",
                "--svg", str(svg_path), "--pptx", str(pptx_in), "--output", str(tmp_path / "out_bad.pptx"),
                "--slide", "0",
                "--trust-origin", "http://localhost:7299",
            ]
            _, out, err, code = run_capturing(ic.main)
        finally:
            sys.argv = old_argv
        check(code not in (None, 0),
              "insert_chart.py: an invalid --trust-origin value is rejected (nonzero exit), never silently ignored")

        # -- CLI-level redaction: a mistakenly-pasted token-bearing URL as ---
        # -- --trust-origin's value must never echo the token to stdout OR --
        # -- stderr, whether from our own message or argparse's own fallback -
        old_argv = sys.argv
        try:
            sys.argv = [
                "insert_chart.py",
                "--svg", str(svg_path), "--pptx", str(pptx_in), "--output", str(tmp_path / "out_bad2.pptx"),
                "--slide", "0",
                "--trust-origin", FIXTURE_LINK,
            ]
            _, out, err, code = run_capturing(ic.main)
        finally:
            sys.argv = old_argv
        check(code not in (None, 0),
              "insert_chart.py: a token-bearing invalid --trust-origin value is rejected")
        check(FIXTURE_TOKEN not in out and FIXTURE_TOKEN not in err,
              "insert_chart.py: the token from a mistakenly-pasted full URL never echoes "
              "to stdout or stderr")


# ---------------------------------------------------------------------------

def main():
    print("Running grunt-pptx script tests...\n")

    print("-- download_chart.py unit tests --")
    test_download_chart_unit()

    print("-- get_svg_info.py unit tests --")
    test_get_svg_info_unit()

    print("-- --trust-origin flag tests (get_svg_info.py + insert_chart.py) --")
    test_trust_origin_flag()

    print("-- download_svg mocked-opener security tests --")
    test_download_svg_security()

    print("-- download_chart.py CLI --trust-origin test --")
    test_download_chart_cli_trust_origin()

    print("-- insert_chart.py functional test --")
    test_insert_chart_functional()

    print("-- insert_chart.py + --trust-origin functional test --")
    test_insert_chart_trust_origin()

    print(f"\n{_passed} check(s) passed, {len(_failures)} failed.")
    if _failures:
        print("\nFAILED:")
        for f in _failures:
            print(f"  - {f}")
        sys.exit(1)
    print("All checks passed.")


if __name__ == "__main__":
    main()
