"""Run with python -m unittest discover -s skills/grunt-pptx/scripts -p test_pptx_quality.py."""
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
import insert_chart
from pptx_validation import validate_ppt_svg, fit_placement, svg_viewport

SVG = b'''<svg xmlns="http://www.w3.org/2000/svg" width="720pt" height="360pt">
<!-- assetId:11111111-1111-1111-1111-111111111111,commitId:aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa -->
<rect width="720" height="360" fill="#234"/><text x="20" y="60" fill="white" font-size="24">EUR millions</text></svg>'''


class PowerPointQualityTests(unittest.TestCase):
    def test_rejects_browser_snapshot_with_export_guidance(self):
        for svg in (SVG.replace(b'width="720pt"', b'data-grunt-scope="table1" width="720pt"'),
                    SVG.replace(b'<rect', b'<style>svg[data-grunt-scope="x"] .t {font-size:52px}</style><rect')):
            with self.subTest(svg=svg):
                with self.assertRaisesRegex(ValueError, 'download_chart.py'):
                    validate_ppt_svg(svg)

    def test_low_level_inserter_cannot_bypass_snapshot_guard(self):
        with self.assertRaisesRegex(ValueError, 'Browser snapshot'):
            insert_chart.add_grunt_web_chart_shape(None, None,
                SVG.replace(b'width="720pt"', b'data-grunt-scope="x" width="720pt"'),
                b'', 0, 0, 1, 1, 'test', {})

    def test_dimensions_respect_units_and_contain_without_distortion(self):
        root = validate_ppt_svg(SVG)
        self.assertEqual((960, 480), svg_viewport(root))
        self.assertEqual((1, 2, 8, 4), fit_placement(root, 1, 1, 8, 6))
        with self.assertRaisesRegex(ValueError, 'distort'):
            fit_placement(root, 1, 1, 8, 6, 'exact')

    def test_rejects_negative_nan_and_infinite_placements(self):
        for values in ((-1, 0, 8, 4), (0, 0, 0, 4), (0, 0, float('nan'), 4), (0, float('inf'), 8, 4)):
            with self.subTest(values=values), self.assertRaises(ValueError):
                fit_placement(validate_ppt_svg(SVG), *values)

    def test_batch_is_atomic_loads_once_and_caches_rasterization(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            deck = directory / 'deck.pptx'
            output = directory / 'out.pptx'
            svg = directory / 'chart.svg'
            svg.write_bytes(SVG)
            presentation = Presentation()
            presentation.slide_width, presentation.slide_height = Inches(12), Inches(7)
            for _ in range(2):
                presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(deck)
            placements = [dict(svg='chart.svg', slide=n, left=1, top=1, width=8, height=4) for n in range(2)]
            # Use the real rasterizer once; inspect the actual saved package, not just function calls.
            with patch.object(insert_chart, 'Presentation', wraps=Presentation) as load, \
                    patch.object(insert_chart, 'svg_to_png', wraps=insert_chart.svg_to_png) as raster:
                result = insert_chart.insert_charts(deck, output, placements, base_dir=directory)
                self.assertEqual(1, load.call_count)
                self.assertEqual(1, raster.call_count)
            self.assertEqual(2, result['inserted'])
            with zipfile.ZipFile(output) as archive:
                svgs = [archive.read(name) for name in archive.namelist() if name.endswith('.svg')]
                self.assertEqual([SVG, SVG], svgs)
            saved = output.read_bytes()
            placements[1]['slide'] = -1
            with self.assertRaisesRegex(ValueError, 'slide index'):
                insert_chart.insert_charts(deck, output, placements, base_dir=directory)
            self.assertEqual(saved, output.read_bytes())

    def test_out_of_bounds_and_unknown_manifest_field_fail_before_save(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            deck, output = directory / 'deck.pptx', directory / 'out.pptx'
            p = Presentation()
            p.slides.add_slide(p.slide_layouts[6])
            p.save(deck)
            (directory / 'chart.svg').write_bytes(SVG)
            item = dict(svg='chart.svg', slide=0, left=8, top=1, width=8, height=4)
            with self.assertRaisesRegex(ValueError, 'beyond'):
                insert_chart.insert_charts(deck, output, [item], base_dir=directory)
            item['left'], item['widht'] = 1, 8
            with self.assertRaisesRegex(ValueError, 'Unknown fields'):
                insert_chart.insert_charts(deck, output, [item], base_dir=directory)
            self.assertFalse(output.exists())

    def test_preview_records_actual_size_and_no_capabilities(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            deck = directory / 'deck.pptx'
            p = Presentation()
            p.slides.add_slide(p.slide_layouts[6])
            p.save(deck)
            (directory / 'chart.svg').write_bytes(SVG)
            result = insert_chart.insert_charts(deck, directory / 'out.pptx',
                [dict(svg='chart.svg', slide=0, left=1, top=1, width=8, height=5)],
                base_dir=directory, preview_dir=directory / 'preview')
            from PIL import Image
            with Image.open(directory / 'preview/placement-001.png') as image:
                self.assertEqual((1152, 576), image.size)
            info = json.loads((directory / 'preview/placements.json').read_text())
            self.assertEqual(4, info['placements'][0]['placement_inches']['height'])
            self.assertNotIn('data-grunt-src', json.dumps(result))


if __name__ == '__main__':
    unittest.main()
