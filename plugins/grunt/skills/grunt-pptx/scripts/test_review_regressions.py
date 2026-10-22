"""Regression cases for export identity, physical sizing and failure atomicity."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from pptx import Presentation

import download_chart
import get_svg_info
import insert_chart
from restore_assets import read_assets
from test_pptx_quality import SVG


class ExportRegressionTests(unittest.TestCase):
    def test_metadata_uses_physical_points_for_all_parsers(self):
        cases = [
            ('width="8in" height="360pt"', (576, 360)),
            ('width="9.6e2px" height="12.7cm"', (720, 360)),
            ('viewBox="0,0,960,480"', (720, 360)),
            ('width="720pt" viewBox="0 0 960 480"', (720, 360)),
        ]
        for parser in (download_chart.parse_svg_metadata, get_svg_info.parse_svg_metadata,
                       insert_chart.parse_svg_metadata):
            for attributes, expected in cases:
                with self.subTest(parser=parser.__module__, attributes=attributes):
                    metadata = parser(f'<svg xmlns="http://www.w3.org/2000/svg" {attributes}/>'.encode())
                    self.assertIsNotNone(metadata['width'])
                    self.assertIsNotNone(metadata['height'])
                    self.assertAlmostEqual(expected[0], metadata['width'])
                    self.assertAlmostEqual(expected[1], metadata['height'])

    def test_rasterizer_refuses_network_and_local_resources(self):
        templates = ('<image href="{0}" width="1" height="1"/>',
                     '<image xmlns:xlink="http://www.w3.org/1999/xlink" xlink:href="{0}" width="1" height="1"/>',
                     '<style>@import url("{0}");</style>',
                     '<rect width="1" height="1" style="fill:url({0})"/>',
                     '<rect width="1" height="1" filter="url(\'{0}\')"/>')
        for resource in ('https://example.invalid/private.png', 'file:///private.png', '/private.png'):
            for template in templates:
                element = template.format(resource)
                svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">{element}</svg>'.encode()
                renderer = MagicMock(side_effect=AssertionError('renderer reached'))
                with self.subTest(element=element), \
                        patch.dict(sys.modules, {'resvg_py': MagicMock(svg_to_bytes=renderer)}), \
                        patch.object(insert_chart, '_cairosvg_png', renderer):
                    with self.assertRaisesRegex(ValueError, 'embedded data URLs'):
                        insert_chart.svg_to_png(svg)
                    renderer.assert_not_called()

    def test_downloader_reports_css_and_attribute_font_families(self):
        svg = b'''<svg xmlns="http://www.w3.org/2000/svg" font-family="Arial" fill="red">
          <style>.a{font-size:9px;font-family:Noto Sans TC, sans-serif;fill:#000}
          .b{font-family:"Segoe UI",sans-serif}</style>
          <text style="font-family: 'Calibri', sans-serif"/><text class="a"/></svg>'''
        self.assertEqual(['Arial', 'Noto Sans TC, sans-serif', 'Segoe UI,sans-serif', 'Calibri, sans-serif'],
                         download_chart.svg_font_families(svg))

    def test_rasterizer_accepts_fragment_and_data_references(self):
        svg = b'''<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">
          <defs><linearGradient id="g"><stop offset="0" stop-color="red"/></linearGradient></defs>
          <style>.a{fill:url(#g)}</style>
          <rect class="a" width="5" height="5"/><rect width="5" height="5" style="fill:url('#g')"/>
          <rect width="5" height="5" fill="url(#g)"/></svg>'''
        insert_chart.require_embedded_resources(svg)

    def test_rasterizer_keeps_embedded_images_and_local_references(self):
        import base64
        import io
        from PIL import Image
        stream = io.BytesIO()
        Image.new('RGB', (1, 1), 'red').save(stream, format='PNG')
        encoded = base64.b64encode(stream.getvalue()).decode()
        svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">
          <defs><rect id="r" width="10" height="10" fill="blue"/></defs><use href="#r"/>
          <image href="data:image/png;base64,{encoded}" width="10" height="10"/></svg>'''.encode()
        with Image.open(io.BytesIO(insert_chart.svg_to_png(svg, 10))) as image:
            red, green, blue = image.convert('RGB').getpixel((5, 5))
            self.assertGreater(red, 200)
            self.assertLess(green, 40)
            self.assertLess(blue, 40)

    def test_failed_preview_leaves_in_place_deck_unchanged(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            deck = self.make_deck(root)
            before = deck.read_bytes()
            (root / 'preview').write_text('not a directory')
            with self.assertRaises(OSError):
                insert_chart.insert_charts(deck, deck, [self.placement()], base_dir=root,
                                           preview_dir=root / 'preview')
            self.assertEqual(before, deck.read_bytes())

    def test_partial_identity_overrides_fail_before_rasterization(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            deck = self.make_deck(root)
            for override in ({'asset_id': '22222222-2222-2222-2222-222222222222'},
                             {'asset_commit': 'b' * 64}):
                with self.subTest(override=override), patch.object(insert_chart, 'svg_to_png', wraps=insert_chart.svg_to_png) as raster:
                    with self.assertRaisesRegex(ValueError, 'supplied together'):
                        insert_chart.insert_charts(deck, root / 'out.pptx',
                                                   [self.placement() | override], base_dir=root)
                    raster.assert_not_called()

    def test_grouped_chart_recovery_rejects_relative_placement(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            deck = self.make_deck(root)
            insert_chart.insert_charts(deck, deck, [self.placement()], base_dir=root)
            prs = Presentation(deck)
            slide = prs.slides[0]
            picture = slide.shapes[0]._element
            group = slide.shapes.add_group_shape()
            group._element.append(picture)
            group._element.recalculate_extents()
            group.left, group.top = 400000, 600000
            group.width *= 2
            prs.save(deck)
            with self.assertRaisesRegex(ValueError, 'Grouped Grunt shape'):
                read_assets(deck)

    def test_private_write_preserves_existing_file_on_replace_failure(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / 'chart.svg'
            destination.write_bytes(b'previous')
            with patch.object(download_chart.os, 'replace', side_effect=OSError('disk failure')):
                with self.assertRaises(OSError):
                    download_chart._write_private_file(str(destination), b'capability')
            self.assertEqual(b'previous', destination.read_bytes())
            self.assertEqual([destination], list(destination.parent.iterdir()))

    @unittest.skipUnless(os.name == 'posix', 'POSIX file modes')
    def test_private_write_replaces_permissive_existing_inode(self):
        import stat
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / 'chart.svg'
            destination.write_bytes(b'previous')
            destination.chmod(0o644)
            download_chart._write_private_file(str(destination), b'capability')
            self.assertEqual(0o600, stat.S_IMODE(destination.stat().st_mode))

    @staticmethod
    def make_deck(root):
        deck = root / 'deck.pptx'
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(deck)
        (root / 'chart.svg').write_bytes(SVG)
        return deck

    @staticmethod
    def placement():
        return dict(svg='chart.svg', slide=0, left=1, top=1, width=8, height=4)


if __name__ == '__main__':
    unittest.main()
